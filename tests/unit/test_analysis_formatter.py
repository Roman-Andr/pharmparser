"""Sheet-level tests for AnalysisFormatter.

The metrics themselves are covered in tests/unit/test_domain_analysis.py; these
check only that the sheet is laid out as expected.
"""

from openpyxl import Workbook

from pharmparser.config import ExportSettings
from pharmparser.domain import PriceTable
from pharmparser.excel.formatters import AnalysisFormatter


def build_rows(settings: ExportSettings, table: PriceTable) -> list[list]:
    ws = Workbook().active
    AnalysisFormatter(settings, table).format(ws)
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_headline_metrics(settings: ExportSettings, table: PriceTable) -> None:
    rows = build_rows(settings, table)
    assert rows[0][0] == "Аптека 1"
    assert rows[1][:2] == ["Асортимент", 3]
    assert rows[2][:2] == ["Средний асортимент конкурентов", 2]
    assert rows[4][:2] == ["Уникальных позиций", 1]


def test_cheapest_everywhere_ignores_missing_items(settings: ExportSettings, table: PriceTable) -> None:
    """Regression for B2, which used to force this metric to 0.

    Аспирин is cheaper at the reference than at either competitor, and Цитрамон is
    stocked only by the reference, so the count is 2. The old float('-inf') sentinel
    made ``price < -inf`` false and produced 0.
    """
    rows = build_rows(settings, table)
    assert rows[3][:2] == ["Позиций ниже всех", 2]


def test_competitor_breakdown(settings: ExportSettings, table: PriceTable) -> None:
    rows = build_rows(settings, table)
    assert rows[5] == ["", "Асортимент", "Дороже", "Дешевле", "Уникальных"]
    assert rows[6] == ["Аптека 2", 2, 1, 1, 0]
    assert rows[7] == ["Аптека 3", 2, 1, 0, 1]


def test_column_widths_are_applied(settings: ExportSettings, table: PriceTable) -> None:
    ws = Workbook().active
    AnalysisFormatter(settings, table).format(ws)
    assert ws.column_dimensions["A"].width == settings.col_width
    assert ws.column_dimensions["B"].width == settings.cell_width
